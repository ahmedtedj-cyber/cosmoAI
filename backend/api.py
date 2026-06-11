#!/usr/bin/env python3
# ═══════════════════════════════════════════════════════════════
#  api.py  —  CosmoAI Backend complet
#  Authentification JWT · SQLite · Analyse Ollama
# ═══════════════════════════════════════════════════════════════
import os, io, csv, json, tempfile
from datetime import datetime, timedelta
from typing import Optional

from fastapi import (
    FastAPI, Depends, HTTPException, UploadFile,
    File, Query, status, BackgroundTasks
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy.orm import Session

from database import create_tables, get_db, User, Product, Analysis
from auth     import (
    hash_password, verify_password, create_access_token,
    get_current_user, ACCESS_TOKEN_EXPIRE_MINUTES
)
from scoring  import compute_health_score, compute_sentiment_ratio, get_top_issue

try:
    from analyzer  import analyze_reviews
    from exporter  import export_json, export_excel
    from config    import MODEL, OLLAMA_URL, EXPORT_FOLDER, DEMO_MODE
    ANALYZER_AVAILABLE = True
except ImportError:
    ANALYZER_AVAILABLE = False
    MODEL       = "non configuré"
    OLLAMA_URL  = "non configuré"
    EXPORT_FOLDER = "exports"
    DEMO_MODE   = True

try:
    from scraper import scrape_reviews
    SCRAPER_AVAILABLE = True
except ImportError:
    SCRAPER_AVAILABLE = False

try:
    from google_places import fetch_google_reviews, search_place
    GOOGLE_PLACES_AVAILABLE = True
except ImportError:
    GOOGLE_PLACES_AVAILABLE = False


# ══════════════════════════════════════════════════════════════
#  App
# ══════════════════════════════════════════════════════════════
app = FastAPI(
    title="CosmoAI API",
    description="Plateforme d'analyse d'avis cosmétiques — Backend complet",
    version="2.0.0",
    docs_url="/docs",
)

# ✅ CORS corrigé — origines explicites sans slash final
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5173",
        "http://127.0.0.1:5174",
        "https://elegant-jelly-32a226.netlify.app",
        "https://cosmoai-app.netlify.app",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def startup():
    create_tables()
    os.makedirs(EXPORT_FOLDER, exist_ok=True)


# ══════════════════════════════════════════════════════════════
#  SCHÉMAS
# ══════════════════════════════════════════════════════════════

class RegisterRequest(BaseModel):
    email:    EmailStr
    username: str = Field(..., min_length=3)
    password: str = Field(..., min_length=6)

class TokenResponse(BaseModel):
    access_token: str
    token_type:   str = "bearer"
    user:         dict

class ProductCreate(BaseModel):
    name:        str = Field(..., min_length=1)
    brand:       str = Field(..., min_length=1)
    category:    Optional[str] = None
    description: Optional[str] = None

class ProductUpdate(BaseModel):
    name:        Optional[str] = None
    brand:       Optional[str] = None
    category:    Optional[str] = None
    description: Optional[str] = None

class AnalyzeRequest(BaseModel):
    product_id: int
    reviews:    list[str] = Field(..., min_length=1)
    verbose:    bool = False


# ══════════════════════════════════════════════════════════════
#  UTILITAIRES
# ══════════════════════════════════════════════════════════════
def _product_or_404(product_id: int, user: User, db: Session) -> Product:
    p = db.query(Product).filter(
        Product.id == product_id,
        Product.owner_id == user.id
    ).first()
    if not p:
        raise HTTPException(404, "Produit introuvable.")
    return p

def _read_csv_bytes(content: bytes, column: str) -> list[str]:
    text   = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    fields = reader.fieldnames or []
    if column not in fields:
        raise HTTPException(422, f"Colonne '{column}' introuvable. Disponibles : {', '.join(fields)}")
    reviews = [row[column].strip() for row in reader if row[column].strip()]
    if not reviews:
        raise HTTPException(422, f"Aucun avis dans '{column}'.")
    return reviews

def _read_excel_bytes(content: bytes, column: str) -> list[str]:
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "openpyxl non installé.")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(content); tmp_path = tmp.name
    try:
        wb = openpyxl.load_workbook(tmp_path)
        ws = wb.worksheets[0]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if column not in headers:
            raise HTTPException(422, f"Colonne '{column}' introuvable. Disponibles : {', '.join(str(h) for h in headers)}")
        ci = headers.index(column)
        reviews = [str(r[ci]).strip() for r in ws.iter_rows(min_row=2, values_only=True)
                   if r[ci] and str(r[ci]).strip()]
    finally:
        os.unlink(tmp_path)
    if not reviews:
        raise HTTPException(422, f"Aucun avis dans '{column}'.")
    return reviews

def _save_analysis(db, user, product_id, reviews, result, source_file=None) -> Analysis:
    health = compute_health_score(result.get("aspects", []), result.get("confidence", "moyenne"))
    analysis = Analysis(
        product_id      = product_id,
        owner_id        = user.id,
        reviews_count   = len(reviews),
        summary         = result.get("summary", ""),
        aspects         = result.get("aspects", []),
        recommendations = result.get("recommendations", []),
        confidence      = result.get("confidence", "moyenne"),
        health_score    = health,
        warnings        = result.get("_warnings", []),
        raw_reviews     = reviews,
        source_file     = source_file,
    )
    db.add(analysis)
    product = db.query(Product).filter(Product.id == product_id).first()
    if product:
        product.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(analysis)
    return analysis

def _format_analysis(a: Analysis) -> dict:
    return {
        "id":             a.id,
        "product_id":     a.product_id,
        "reviews_count":  a.reviews_count,
        "summary":        a.summary,
        "aspects":        a.aspects,
        "recommendations":a.recommendations,
        "confidence":     a.confidence,
        "health_score":   a.health_score,
        "sentiment_ratio":compute_sentiment_ratio(a.aspects or []),
        "top_issue":      get_top_issue(a.aspects or []),
        "_warnings":      a.warnings,
        "source_file":    a.source_file,
        "created_at":     a.created_at.isoformat(),
    }


# ══════════════════════════════════════════════════════════════
#  ROUTES
# ══════════════════════════════════════════════════════════════

@app.get("/status", tags=["Système"])
def get_status(current_user: User = Depends(get_current_user)):
    return {
        "status":         "ok",
        "model":          MODEL if not DEMO_MODE else "demo (mock)",
        "demo_mode":      DEMO_MODE,
        "analyzer_ready": ANALYZER_AVAILABLE,
        "timestamp":      datetime.utcnow().isoformat(),
    }


# ══════════════════════════════════════════════════════════════
#  AUTH
# ══════════════════════════════════════════════════════════════

@app.post("/auth/register", tags=["Auth"], status_code=201)
def register(payload: RegisterRequest, db: Session = Depends(get_db)):
    if db.query(User).filter(User.email == payload.email).first():
        raise HTTPException(400, "Email déjà utilisé.")
    if db.query(User).filter(User.username == payload.username).first():
        raise HTTPException(400, "Nom d'utilisateur déjà pris.")
    user = User(
        email           = payload.email,
        username        = payload.username,
        hashed_password = hash_password(payload.password[:72]),
    )
    db.add(user); db.commit(); db.refresh(user)
    token = create_access_token({"sub": str(user.id)})
    return {
        "access_token": token,
        "token_type":   "bearer",
        "user": {"id": user.id, "email": user.email, "username": user.username},
    }


@app.post("/auth/login", tags=["Auth"], response_model=TokenResponse)
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = (
        db.query(User).filter(User.email == form.username).first() or
        db.query(User).filter(User.username == form.username).first()
    )
    if not user or not verify_password(form.password[:72], user.hashed_password):
        raise HTTPException(401, "Identifiants incorrects.")
    if not user.is_active:
        raise HTTPException(403, "Compte désactivé.")
    token = create_access_token({"sub": str(user.id)})
    return {
        "access_token": token,
        "token_type":   "bearer",
        "user": {"id": user.id, "email": user.email, "username": user.username},
    }


@app.get("/auth/me", tags=["Auth"])
def me(current_user: User = Depends(get_current_user)):
    return {
        "id":         current_user.id,
        "email":      current_user.email,
        "username":   current_user.username,
        "created_at": current_user.created_at.isoformat(),
    }


# ══════════════════════════════════════════════════════════════
#  PRODUCTS
# ══════════════════════════════════════════════════════════════

@app.get("/products", tags=["Produits"])
def list_products(
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    products = db.query(Product).filter(Product.owner_id == current_user.id).all()
    result = []
    for p in products:
        latest = (
            db.query(Analysis)
            .filter(Analysis.product_id == p.id)
            .order_by(Analysis.created_at.desc())
            .first()
        )
        result.append({
            "id":           p.id,
            "name":         p.name,
            "brand":        p.brand,
            "category":     p.category,
            "description":  p.description,
            "total_reviews":sum(a.reviews_count for a in p.analyses),
            "health_score": latest.health_score if latest else None,
            "top_issue":    get_top_issue(latest.aspects or []) if latest else None,
            "last_analysis":latest.created_at.isoformat() if latest else None,
            "updated_at":   p.updated_at.isoformat(),
            "created_at":   p.created_at.isoformat(),
        })
    return {"products": result, "count": len(result)}


@app.post("/products", tags=["Produits"], status_code=201)
def create_product(
    payload:      ProductCreate,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    product = Product(**payload.dict(), owner_id=current_user.id)
    db.add(product); db.commit(); db.refresh(product)
    return {"id": product.id, "name": product.name, "brand": product.brand,
            "category": product.category, "description": product.description,
            "created_at": product.created_at.isoformat()}


@app.get("/products/{product_id}", tags=["Produits"])
def get_product(
    product_id:   int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    p = _product_or_404(product_id, current_user, db)
    analyses = db.query(Analysis).filter(Analysis.product_id == p.id)\
                 .order_by(Analysis.created_at.desc()).all()
    trend = [
        {"date": a.created_at.strftime("%Y-%m-%d"), "health_score": a.health_score,
         "reviews_count": a.reviews_count}
        for a in analyses
    ]
    latest = analyses[0] if analyses else None
    return {
        "id": p.id, "name": p.name, "brand": p.brand,
        "category": p.category, "description": p.description,
        "total_analyses": len(analyses),
        "total_reviews":  sum(a.reviews_count for a in analyses),
        "health_score":   latest.health_score if latest else None,
        "top_issue":      get_top_issue(latest.aspects or []) if latest else None,
        "sentiment_ratio":compute_sentiment_ratio(latest.aspects or []) if latest else None,
        "trend":          trend,
        "created_at":     p.created_at.isoformat(),
        "updated_at":     p.updated_at.isoformat(),
    }


@app.put("/products/{product_id}", tags=["Produits"])
def update_product(
    product_id:   int,
    payload:      ProductUpdate,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    p = _product_or_404(product_id, current_user, db)
    for field, val in payload.dict(exclude_none=True).items():
        setattr(p, field, val)
    p.updated_at = datetime.utcnow()
    db.commit(); db.refresh(p)
    return {"id": p.id, "name": p.name, "brand": p.brand,
            "category": p.category, "updated_at": p.updated_at.isoformat()}


@app.delete("/products/{product_id}", tags=["Produits"], status_code=204)
def delete_product(
    product_id:   int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    p = _product_or_404(product_id, current_user, db)
    db.delete(p); db.commit()


# ══════════════════════════════════════════════════════════════
#  ANALYSE
# ══════════════════════════════════════════════════════════════

@app.post("/analyze", tags=["Analyse"])
def analyze(
    payload:      AnalyzeRequest,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Moteur d'analyse non disponible (analyzer.py introuvable).")
    _product_or_404(payload.product_id, current_user, db)
    try:
        result = analyze_reviews(payload.reviews, verbose=payload.verbose)
    except ConnectionError as e:
        raise HTTPException(503, str(e))
    except TimeoutError as e:
        raise HTTPException(504, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    analysis = _save_analysis(db, current_user, payload.product_id, payload.reviews, result)
    return _format_analysis(analysis)


@app.post("/analyze/csv", tags=["Analyse"])
async def analyze_csv(
    product_id:   int        = Query(...),
    file:         UploadFile = File(...),
    column:       str        = Query("avis"),
    verbose:      bool       = Query(False),
    current_user: User       = Depends(get_current_user),
    db:           Session    = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Moteur d'analyse non disponible.")
    _product_or_404(product_id, current_user, db)
    content = await file.read()
    reviews = _read_csv_bytes(content, column)
    try:
        result = analyze_reviews(reviews, verbose=verbose)
    except ConnectionError as e:
        raise HTTPException(503, str(e))
    except TimeoutError as e:
        raise HTTPException(504, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    analysis = _save_analysis(db, current_user, product_id, reviews, result, file.filename)
    return _format_analysis(analysis)


@app.post("/analyze/excel", tags=["Analyse"])
async def analyze_excel(
    product_id:   int        = Query(...),
    file:         UploadFile = File(...),
    column:       str        = Query("avis"),
    verbose:      bool       = Query(False),
    current_user: User       = Depends(get_current_user),
    db:           Session    = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Moteur d'analyse non disponible.")
    _product_or_404(product_id, current_user, db)
    content = await file.read()
    reviews = _read_excel_bytes(content, column)
    try:
        result = analyze_reviews(reviews, verbose=verbose)
    except ConnectionError as e:
        raise HTTPException(503, str(e))
    except TimeoutError as e:
        raise HTTPException(504, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    analysis = _save_analysis(db, current_user, product_id, reviews, result, file.filename)
    return _format_analysis(analysis)


# ══════════════════════════════════════════════════════════════
#  SCRAPING URL
# ══════════════════════════════════════════════════════════════

class ScrapeRequest(BaseModel):
    product_id: int
    url:        str
    verbose:    bool = False

@app.post("/analyze/url", tags=["Analyse"])
def analyze_url(
    payload:      ScrapeRequest,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Moteur d'analyse non disponible.")
    if not SCRAPER_AVAILABLE:
        raise HTTPException(503, "Module scraper.py introuvable.")
    _product_or_404(payload.product_id, current_user, db)
    try:
        reviews, meta = scrape_reviews(payload.url)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except ConnectionError as e:
        raise HTTPException(503, f"Site inaccessible : {e}")
    except RuntimeError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(500, f"Erreur de scraping : {e}")
    try:
        result = analyze_reviews(reviews, verbose=payload.verbose)
    except ConnectionError as e:
        raise HTTPException(503, str(e))
    except TimeoutError as e:
        raise HTTPException(504, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    source_label = f"url:{meta.get('source','web')}:{payload.url[:80]}"
    analysis = _save_analysis(db, current_user, payload.product_id, reviews, result, source_label)
    resp = _format_analysis(analysis)
    resp["scrape_meta"] = {
        "reviews_found": len(reviews),
        "source":        meta.get("source"),
        "product_hint":  meta.get("product_hint", ""),
        "pages_scraped": meta.get("pages_scraped", 1),
    }
    return resp


@app.get("/scrape/preview", tags=["Analyse"])
def scrape_preview(
    url:          str  = Query(...),
    current_user: User = Depends(get_current_user),
):
    if not SCRAPER_AVAILABLE:
        raise HTTPException(503, "Module scraper.py introuvable.")
    try:
        reviews, meta = scrape_reviews(url)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except ConnectionError as e:
        raise HTTPException(503, f"Site inaccessible : {e}")
    except RuntimeError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(500, f"Erreur de scraping : {e}")
    return {
        "reviews":       reviews[:5],
        "total":         len(reviews),
        "source":        meta.get("source"),
        "product_hint":  meta.get("product_hint", ""),
        "pages_scraped": meta.get("pages_scraped", 1),
    }


# ══════════════════════════════════════════════════════════════
#  GOOGLE PLACES
# ══════════════════════════════════════════════════════════════

class GoogleSearchRequest(BaseModel):
    product_id: int
    query:      str
    api_key:    Optional[str] = None

@app.get("/google/search", tags=["Google Places"])
def google_search(
    query:        str  = Query(...),
    api_key:      Optional[str] = Query(None),
    current_user: User = Depends(get_current_user),
):
    if not GOOGLE_PLACES_AVAILABLE:
        raise HTTPException(503, "Module google_places.py introuvable.")
    key = api_key or os.getenv("GOOGLE_PLACES_API_KEY", "")
    if not key:
        raise HTTPException(400, "Clé API Google manquante.")
    try:
        place = search_place(query, api_key=key)
        return place
    except RuntimeError as e:
        raise HTTPException(404, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))


@app.post("/analyze/google", tags=["Google Places"])
def analyze_google(
    payload:      GoogleSearchRequest,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    if not GOOGLE_PLACES_AVAILABLE:
        raise HTTPException(503, "Module google_places.py introuvable.")
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Moteur d'analyse non disponible.")
    _product_or_404(payload.product_id, current_user, db)
    key = payload.api_key or os.getenv("GOOGLE_PLACES_API_KEY", "")
    if not key:
        raise HTTPException(400, "Clé API Google manquante.")
    try:
        reviews, meta = fetch_google_reviews(payload.query, api_key=key)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except RuntimeError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    try:
        result = analyze_reviews(reviews)
    except ConnectionError as e:
        raise HTTPException(503, str(e))
    except TimeoutError as e:
        raise HTTPException(504, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))
    source_label = f"google:{meta.get('place_id','')}:{meta.get('name','')[:60]}"
    analysis = _save_analysis(db, current_user, payload.product_id, reviews, result, source_label)
    resp = _format_analysis(analysis)
    resp["google_meta"] = meta
    return resp


# ══════════════════════════════════════════════════════════════
#  HISTORIQUE
# ══════════════════════════════════════════════════════════════

@app.get("/history", tags=["Historique"])
def list_history(
    product_id:   Optional[int] = Query(None),
    limit:        int           = Query(50, le=200),
    current_user: User          = Depends(get_current_user),
    db:           Session       = Depends(get_db),
):
    q = db.query(Analysis).filter(Analysis.owner_id == current_user.id)
    if product_id:
        q = q.filter(Analysis.product_id == product_id)
    analyses = q.order_by(Analysis.created_at.desc()).limit(limit).all()
    return {
        "analyses": [
            {
                "id":            a.id,
                "product_id":    a.product_id,
                "product_name":  a.product.name,
                "reviews_count": a.reviews_count,
                "health_score":  a.health_score,
                "confidence":    a.confidence,
                "top_issue":     get_top_issue(a.aspects or []),
                "source_file":   a.source_file,
                "created_at":    a.created_at.isoformat(),
            }
            for a in analyses
        ],
        "count": len(analyses),
    }


@app.get("/history/{analysis_id}", tags=["Historique"])
def get_analysis(
    analysis_id:  int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    a = db.query(Analysis).filter(
        Analysis.id == analysis_id,
        Analysis.owner_id == current_user.id
    ).first()
    if not a:
        raise HTTPException(404, "Analyse introuvable.")
    return _format_analysis(a)


@app.delete("/history/{analysis_id}", tags=["Historique"], status_code=204)
def delete_analysis(
    analysis_id:  int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    a = db.query(Analysis).filter(
        Analysis.id == analysis_id,
        Analysis.owner_id == current_user.id
    ).first()
    if not a:
        raise HTTPException(404, "Analyse introuvable.")
    db.delete(a); db.commit()


# ══════════════════════════════════════════════════════════════
#  DASHBOARD
# ══════════════════════════════════════════════════════════════

@app.get("/dashboard", tags=["Dashboard"])
def dashboard(
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    analyses = db.query(Analysis).filter(Analysis.owner_id == current_user.id).all()
    products = db.query(Product).filter(Product.owner_id == current_user.id).all()

    total_reviews  = sum(a.reviews_count for a in analyses)
    total_analyses = len(analyses)

    all_aspects = [asp for a in analyses for asp in (a.aspects or [])]
    ratio     = compute_sentiment_ratio(all_aspects)
    top_issue = get_top_issue(all_aspects)

    scores    = [a.health_score for a in analyses if a.health_score is not None]
    avg_score = round(sum(scores) / len(scores), 1) if scores else 0.0

    cutoff = datetime.utcnow() - timedelta(days=30)
    recent = [a for a in analyses if a.created_at >= cutoff]
    trend  = {}
    for a in recent:
        day = a.created_at.strftime("%Y-%m-%d")
        trend.setdefault(day, {"date": day, "analyses": 0, "reviews": 0})
        trend[day]["analyses"] += 1
        trend[day]["reviews"]  += a.reviews_count
    trend_list = sorted(trend.values(), key=lambda x: x["date"])

    insights = []
    neg_aspects = sorted(
        [a for a in all_aspects if a.get("sentiment") == "négatif"],
        key=lambda x: x.get("mentions", 0), reverse=True
    )
    pos_aspects = sorted(
        [a for a in all_aspects if a.get("sentiment") == "positif"],
        key=lambda x: x.get("mentions", 0), reverse=True
    )
    if neg_aspects:
        insights.append({"type": "warning", "text": f"Problème principal : {neg_aspects[0]['nom']} ({neg_aspects[0].get('mentions',0)} mentions)"})
    if pos_aspects:
        insights.append({"type": "success", "text": f"Point fort : {pos_aspects[0]['nom']} ({pos_aspects[0].get('mentions',0)} mentions)"})
    if ratio.get("négatif", 0) > 40:
        insights.append({"type": "alert", "text": f"{ratio['négatif']}% des aspects sont négatifs — action requise."})
    elif ratio.get("positif", 0) > 60:
        insights.append({"type": "info", "text": f"Sentiment globalement positif ({ratio['positif']}%)."})

    return {
        "kpis": {
            "total_reviews":    total_reviews,
            "total_analyses":   total_analyses,
            "total_products":   len(products),
            "avg_health_score": avg_score,
            "sentiment_ratio":  ratio,
            "top_issue":        top_issue,
        },
        "insights":  insights,
        "trend_30d": trend_list,
        "recent_analyses": [
            {
                "id":            a.id,
                "product_name":  a.product.name,
                "health_score":  a.health_score,
                "reviews_count": a.reviews_count,
                "created_at":    a.created_at.isoformat(),
            }
            for a in sorted(analyses, key=lambda x: x.created_at, reverse=True)[:5]
        ],
    }


# ══════════════════════════════════════════════════════════════
#  COMPARAISON
# ══════════════════════════════════════════════════════════════

@app.get("/compare", tags=["Comparaison"])
def compare_products(
    product_a:    int     = Query(...),
    product_b:    int     = Query(...),
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    pa = _product_or_404(product_a, current_user, db)
    pb = _product_or_404(product_b, current_user, db)

    def _product_summary(p: Product):
        analyses = db.query(Analysis).filter(Analysis.product_id == p.id)\
                     .order_by(Analysis.created_at.desc()).all()
        if not analyses:
            return {"id": p.id, "name": p.name, "brand": p.brand,
                    "health_score": None, "sentiment_ratio": None,
                    "top_aspects": [], "top_issue": None, "total_reviews": 0}
        all_asp = [asp for a in analyses for asp in (a.aspects or [])]
        latest  = analyses[0]
        return {
            "id":              p.id,
            "name":            p.name,
            "brand":           p.brand,
            "health_score":    latest.health_score,
            "sentiment_ratio": compute_sentiment_ratio(all_asp),
            "top_aspects":     sorted(all_asp, key=lambda x: x.get("mentions", 0), reverse=True)[:5],
            "top_issue":       get_top_issue(all_asp),
            "total_reviews":   sum(a.reviews_count for a in analyses),
            "total_analyses":  len(analyses),
        }

    sa = _product_summary(pa)
    sb = _product_summary(pb)
    differences = []
    if sa["health_score"] and sb["health_score"]:
        diff   = sa["health_score"] - sb["health_score"]
        better = sa["name"] if diff > 0 else sb["name"]
        differences.append(f"{better} a un meilleur score de santé ({abs(diff):.1f} pts d'écart).")

    return {"product_a": sa, "product_b": sb, "key_differences": differences}


# ══════════════════════════════════════════════════════════════
#  EXPORT
# ══════════════════════════════════════════════════════════════

@app.get("/export/{analysis_id}/json", tags=["Export"])
def export_analysis_json(
    analysis_id:  int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Export non disponible.")
    a = db.query(Analysis).filter(
        Analysis.id == analysis_id, Analysis.owner_id == current_user.id
    ).first()
    if not a:
        raise HTTPException(404, "Analyse introuvable.")
    path = export_json(_format_analysis(a), f"analyse_{analysis_id}.json")
    return FileResponse(path, media_type="application/json",
                        filename=f"cosmoai_analyse_{analysis_id}.json")


@app.get("/export/{analysis_id}/excel", tags=["Export"])
def export_analysis_excel(
    analysis_id:  int,
    current_user: User    = Depends(get_current_user),
    db:           Session = Depends(get_db),
):
    if not ANALYZER_AVAILABLE:
        raise HTTPException(503, "Export non disponible.")
    a = db.query(Analysis).filter(
        Analysis.id == analysis_id, Analysis.owner_id == current_user.id
    ).first()
    if not a:
        raise HTTPException(404, "Analyse introuvable.")
    path = export_excel(_format_analysis(a), f"analyse_{analysis_id}.xlsx")
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"cosmoai_analyse_{analysis_id}.xlsx"
    )