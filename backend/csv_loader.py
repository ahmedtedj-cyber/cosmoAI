# ─────────────────────────────────────────
#  csv_loader.py  —  Chargement des avis
#  Supporte : CSV, Excel (.xlsx)
# ─────────────────────────────────────────
import csv
import os


def load_reviews_from_csv(filepath: str, column: str = "avis") -> list[str]:
    """
    Charge les avis depuis un fichier CSV.

    Args:
        filepath : chemin vers le fichier .csv
        column   : nom de la colonne contenant les avis (défaut: "avis")

    Returns:
        liste de strings (avis nettoyés)
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Fichier introuvable : {filepath}")

    reviews = []
    with open(filepath, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)

        # Vérifie que la colonne existe
        if column not in reader.fieldnames:
            available = ", ".join(reader.fieldnames)
            raise ValueError(
                f"Colonne '{column}' introuvable dans le CSV.\n"
                f"Colonnes disponibles : {available}"
            )

        for row in reader:
            text = row[column].strip()
            if text:   # ignore les lignes vides
                reviews.append(text)

    if not reviews:
        raise ValueError(f"Aucun avis trouvé dans la colonne '{column}'.")

    return reviews


def load_reviews_from_excel(filepath: str, column: str = "avis", sheet: int = 0) -> list[str]:
    """
    Charge les avis depuis un fichier Excel (.xlsx).
    Nécessite : pip install openpyxl

    Args:
        filepath : chemin vers le fichier .xlsx
        column   : nom de la colonne contenant les avis
        sheet    : index de la feuille (0 = première)
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("Installe openpyxl : pip install openpyxl")

    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Fichier introuvable : {filepath}")

    wb      = openpyxl.load_workbook(filepath)
    ws      = wb.worksheets[sheet]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    if column not in headers:
        raise ValueError(
            f"Colonne '{column}' introuvable dans le fichier Excel.\n"
            f"Colonnes disponibles : {', '.join(str(h) for h in headers)}"
        )

    col_idx = headers.index(column)
    reviews = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[col_idx]
        if val and str(val).strip():
            reviews.append(str(val).strip())

    if not reviews:
        raise ValueError(f"Aucun avis trouvé dans la colonne '{column}'.")

    return reviews
