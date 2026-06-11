# ─────────────────────────────────────────
#  exporter.py  —  Export des résultats
#  Formats : JSON, Excel
# ─────────────────────────────────────────
import json
import os
from datetime import datetime
from config import EXPORT_FOLDER


def _ensure_folder():
    os.makedirs(EXPORT_FOLDER, exist_ok=True)


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


# ── JSON ────────────────────────────────────────────────────────────────
def export_json(result: dict, filename: str = None) -> str:
    """Sauvegarde le résultat brut en JSON."""
    _ensure_folder()
    filename = filename or f"analyse_{_timestamp()}.json"
    path     = os.path.join(EXPORT_FOLDER, filename)

    with open(path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    return path


# ── Excel ───────────────────────────────────────────────────────────────
def export_excel(result: dict, filename: str = None) -> str:
    """
    Exporte le résultat dans un fichier Excel .xlsx
    avec 3 feuilles : Aspects | Recommandations | Résumé
    Nécessite : pip install openpyxl
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("Installe openpyxl : pip install openpyxl")

    _ensure_folder()
    filename = filename or f"analyse_{_timestamp()}.xlsx"
    path     = os.path.join(EXPORT_FOLDER, filename)

    wb = openpyxl.Workbook()

    # ── Feuille 1 : Aspects ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Aspects"

    header_fill = PatternFill("solid", fgColor="1E1E2E")
    header_font = Font(bold=True, color="FFFFFF")

    headers = ["Aspect", "Sentiment", "Mentions", "Exemples"]
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    sentiment_colors = {
        "positif": "C6EFCE",
        "négatif": "FFC7CE",
        "mitigé":  "FFEB9C",
    }

    for row, aspect in enumerate(result.get("aspects", []), 2):
        ws1.cell(row=row, column=1, value=aspect.get("nom", ""))
        s_cell = ws1.cell(row=row, column=2, value=aspect.get("sentiment", ""))
        color  = sentiment_colors.get(aspect.get("sentiment", ""), "FFFFFF")
        s_cell.fill = PatternFill("solid", fgColor=color)
        ws1.cell(row=row, column=3, value=aspect.get("mentions", 0))
        exemples = " | ".join(aspect.get("exemples", []))
        ws1.cell(row=row, column=4, value=exemples)

    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 30

    # ── Feuille 2 : Recommandations ──────────────────────────────────
    ws2 = wb.create_sheet("Recommandations")
    ws2.cell(row=1, column=1, value="Recommandations").font = Font(bold=True, size=13)
    for i, rec in enumerate(result.get("recommendations", []), 2):
        ws2.cell(row=i, column=1, value=f"{i-1}. {rec}")
    ws2.column_dimensions["A"].width = 80

    # ── Feuille 3 : Résumé ───────────────────────────────────────────
    ws3 = wb.create_sheet("Résumé")
    data = [
        ("Résumé",          result.get("summary", "")),
        ("Confiance IA",    result.get("confidence", "")),
        ("Avis analysés",   result.get("avis_analysés", 0)),
        ("Alertes validation", "\n".join(result.get("_warnings", []))),
    ]
    for row, (label, value) in enumerate(data, 1):
        ws3.cell(row=row, column=1, value=label).font = Font(bold=True)
        cell = ws3.cell(row=row, column=2, value=str(value))
        cell.alignment = Alignment(wrap_text=True)

    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 80

    wb.save(path)
    return path
